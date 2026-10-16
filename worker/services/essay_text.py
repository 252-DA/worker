"""Lấy chữ từ file bài luận sinh viên nộp qua Canvas để chấm.

Chỉ cần văn bản thuần, không cần cấu trúc heading như khi đánh chỉ mục tài
liệu, nên đọc thẳng bằng PyMuPDF/python-docx thay vì qua parser của pipeline
(parser PDF thích ứng có thể chạy OCR rất lâu). Bài nhập chữ trực tiếp trong
Canvas đến dưới dạng HTML.
"""
from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".html", ".htm", ".txt", ".md")


class UnsupportedAsset(ValueError):
    """File không đọc được chữ: định dạng lạ hoặc PDF chỉ có ảnh."""


class _TextOnly(HTMLParser):
    _BLOCKS = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in self._BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data):
        self.parts.append(data)


def _html(raw: str) -> str:
    parser = _TextOnly()
    parser.feed(raw)
    return "".join(parser.parts)


def _pdf(path: Path) -> str:
    import fitz

    with fitz.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _clean(text: str) -> str:
    lines = (line.strip() for line in text.splitlines())
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()


def extract_text(path: Path, content_type: str | None = None) -> str:
    """Văn bản của một file; ``UnsupportedAsset`` khi không đọc được chữ nào."""
    ext = path.suffix.lower()
    kind = (content_type or "").split(";")[0].strip().lower()
    if ext == ".pdf" or kind == "application/pdf":
        text = _pdf(path)
    elif ext == ".docx":
        text = _docx(path)
    elif ext in (".html", ".htm") or kind == "text/html":
        text = _html(path.read_text(encoding="utf-8", errors="replace"))
    elif ext in (".txt", ".md") or kind.startswith("text/"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise UnsupportedAsset(
            f"Không đọc được file {path.name}: chỉ nhận {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    text = _clean(text)
    if not text:
        raise UnsupportedAsset(
            f"File {path.name} không có chữ đọc được (có thể là PDF scan chỉ có ảnh)"
        )
    return text


__all__ = ["SUPPORTED_EXTENSIONS", "UnsupportedAsset", "extract_text"]
